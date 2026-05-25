# -*- coding: utf-8 -*-
"""根据简约版简历生成含 AI 个人项目的 Word 简历。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

OUTPUT = Path(r"c:\Users\yoeec\Desktop\docuemnt\陈永义简历简约版.docx")


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=12, bold=False)


def add_section(doc, title):
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_run_font(run, size=13, bold=True, color=(0, 70, 140))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=bold)
    p.paragraph_format.space_after = Pt(2)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(1)


def build():
    doc = Document()
    # 默认正文字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    add_title(doc, "个人简历")
    add_subtitle(doc, "陈永义｜Java / 全栈 / AI 开发工程师")
    add_body(doc, "电话：13002860228    邮箱：yoee.chan.yy@gmail.com    所在地：成都")
    add_body(doc, "工作年限：4 年    学历：本科（西昌学院 · 计算机科学与技术 · 2022 年毕业）")
    add_body(doc, "求职方向：Java 中级工程师 / 全栈工程师 / AI 开发工程师")

    add_section(doc, "掌握技能")
    add_body(doc, "后端开发（Java）", bold=True)
    for line in [
        "Java 语言体系：集合、IO/NIO、反射、并发模型、锁语义",
        "JVM：类加载机制、JMM、逃逸分析、栈上分配、CMS/G1/ZGC、GC 调优",
        "Spring 生态：Spring Boot、MyBatis、AOP、拦截器链、Bean 生命周期",
        "设计能力：DDD 领域建模、策略模式、责任链、模板方法",
    ]:
        add_bullet(doc, line)

    add_body(doc, "分布式系统与中间件", bold=True)
    add_bullet(doc, "Kafka：producer 调优、log segment、事务语义、幂等生产、位移提交、积压排查")
    add_bullet(doc, "Redis：数据结构、过期策略、分布式锁")
    add_bullet(doc, "SQL Server / MySQL：索引优化、慢 SQL 分析、表结构设计")

    add_body(doc, "AI 工程与 Python", bold=True)
    for line in [
        "Transformer、RAG、LangChain / LangChain4j、Prompt 模板",
        "规则 + LLM 混合任务流、结构化抽取",
        "Agent 架构、MCP 协议、本地 LLM 部署、记忆压缩（EMC）",
        "pandas/numpy、大文件处理、Excel 清洗与 DB 入库",
    ]:
        add_bullet(doc, line)

    add_body(doc, "前端与工程化", bold=True)
    for line in [
        "Vue / React / uni-app / Electron",
        "Docker、K8S、CI/CD（Jenkins/GitLab CI）、灰度发布",
        "Git、Maven、Postman、Jira、Copilot",
    ]:
        add_bullet(doc, line)

    add_section(doc, "技能证书")
    add_bullet(doc, "获取微软专家徽章（微软云原生开发）")
    add_bullet(doc, "大学英语四级，良好的听说读写能力，快速浏览英语专业文件及书籍。")

    add_section(doc, "自我评价")
    for line in [
        "强 Owner 意识，能独立推进复杂任务",
        "具备产品思维，能在业务价值与技术可行性之间做平衡",
        "沟通能力强，跨团队协作经验丰富",
        "自驱力强，持续学习，擅长解决复杂技术问题",
        "喜欢阅读历史文化类作品，性格稳重、专注",
    ]:
        add_bullet(doc, line)

    add_section(doc, "个人项目")
    add_body(doc, "KongMing（Yolk）— 本地化 AI Agent 桌面应用（2025 – 至今）", bold=True)
    add_body(
        doc,
        "技术：Electron、React、TypeScript、Python、LangChain/OpenManus、MCP、Docker、"
        "结构化配置（JSON/TOML）、IPC 流式通信",
    )
    add_body(doc, "项目简介", bold=True)
    add_bullet(
        doc,
        "自研跨平台桌面 Agent：Electron 负责系统能力与 UI，Python 负责 Agent 执行与工具调用，"
        "支持本地/云端 LLM 一键切换，可打包为 Windows 安装包实现本地化部署。",
    )
    add_bullet(
        doc,
        "双 Agent 架构：流式对话（agent_stream）+ 一次性工具 API（agent_tools），"
        "支持 need_input 人机协同、多步骤任务规划与执行。",
    )
    add_body(doc, "核心职责与技术亮点", bold=True)
    for line in [
        "设计 Renderer → Electron Main → Python Agent → LLM 四层架构及 IPC 通道",
        "实现 Setting 后端：Factory + Repository + Validator 可扩展配置体系（MCP/LLM/Workspace）",
        "MCP 服务器可配置、工作区读写权限、软件操作权限与三级风险确认策略",
        "Episodic Memory Compression（EMC）：按时间/主题/线索的结构化记忆压缩与回忆重建",
        "微信公众号自动发布流水线：OAuth、草稿创建、Skill 编排与连接器集成",
        "工具能力：文件读写、Excel 聚合提取、Jira/邮件等连接器、流水线任务与词元流量管理 UI",
    ]:
        add_bullet(doc, line)
    add_body(doc, "项目成果", bold=True)
    for line in [
        "完成 MVP：流式对话、工具调用、配置分离、安装包打包链路",
        "形成可复用的 Agent + MCP + 记忆 + 权限模型，支撑后续插件化与商业化扩展",
    ]:
        add_bullet(doc, line)

    add_section(doc, "重点项目介绍")
    add_body(doc, "公司：友邦创新科技（2022 年 8 月 – 至今）", bold=True)

    add_body(doc, "AML 服务中间件（2024.08 – 至今）", bold=True)
    add_body(doc, "技术：Spring Boot、Kafka、Redis、MyBatis、K8S、Docker、SFTP、Blob Storage")
    add_body(doc, "项目简介", bold=True)
    add_bullet(
        doc,
        "为保险业务提供黑名单校验中间服务，负责上下游系统集成、数据解密、文件扫描与高并发查询。",
    )
    add_body(doc, "职责", bold=True)
    for line in [
        "负责 Kafka 消息处理链路（反序列化、解密、异步入库）",
        "设计线程池、连接池、缓存策略；开发本地缓存",
        "参与需求设计、迭代开发与上线；流程优化和系统迭代开发",
    ]:
        add_bullet(doc, line)
    add_body(doc, "技术难点与解决方案", bold=True)
    for line in [
        "高并发下 LRU + sync 锁实现线程安全本地缓存",
        "扫描文件使用内存池降低 GC 压力；线程池与 HikariCP 参数匹配调优",
        "调整 JVM 栈大小解决 XML 回溯 StackOverflow；binlog + ES 定位慢 SQL",
        "消费端幂等性与数据一致性，优化分区策略与消费速率",
    ]:
        add_bullet(doc, line)
    add_body(doc, "成果", bold=True)
    for line in [
        "系统吞吐量提升 1.5 倍；查询延迟降低 20%",
        "Kafka 消息处理稳定性显著提升",
    ]:
        add_bullet(doc, line)

    add_body(doc, "S8 自动标注系统（2025.09 – 至今）", bold=True)
    add_body(doc, "技术：Python、Java、Agent、Spring Boot、K8S、Docker、SQL Server")
    add_body(doc, "项目简介", bold=True)
    add_bullet(
        doc,
        "在原有订单系统中引入 AI，实现基于姓名、日期、国家等 8 类规则的自动标注，提高订单自动化处理率。",
    )
    add_body(doc, "职责", bold=True)
    for line in [
        "负责 AI Agent 的概念验证与核心开发",
        "设计「Agent → 结构化抽取 → DB → Java 规则引擎」架构",
        "负责日期解析、人名拆分、国家标准化等难抽取字段；参与 Java 与 Agent 系统集成",
    ]:
        add_bullet(doc, line)
    add_body(doc, "技术亮点", bold=True)
    for line in [
        "规则 + LLM 混合任务流；Prompt 模板 + JSON 结构化输出",
        "Python 负责抽取，Java 负责规则与流程；AI 模块可复用、可扩展",
    ]:
        add_bullet(doc, line)
    add_body(doc, "成果", bold=True)
    for line in [
        "自动标注率显著提升，大幅减少人工审核成本",
        "AI 模块成功集成到原系统",
    ]:
        add_bullet(doc, line)

    add_body(doc, "Cashflow Transfer Portal（2023.04 – 2024.08）", bold=True)
    add_body(doc, "技术：Python、React、SQL Server、SFTP")
    add_body(doc, "项目简介", bold=True)
    add_bullet(
        doc,
        "跨部门协作的现金流交易系统，从 0 到 1 落地，负责数据清洗、汇率换算、INC 文件生成与 SFTP 上传到彭博社。",
    )
    add_body(doc, "职责", bold=True)
    for line in [
        "核心技术设计与开发；Python 处理 Excel 数据清洗、聚合、分类、入库",
        "设计环队列算法生成 SFTP 文件；参与跨部门需求沟通与方案评审",
    ]:
        add_bullet(doc, line)
    add_body(doc, "成果", bold=True)
    for line in [
        "项目成功上线，打通多团队业务链路",
        "显著提升现金流处理自动化程度，为团队带来新业务增长",
    ]:
        add_bullet(doc, line)

    add_section(doc, "实习项目介绍")
    add_body(doc, "烟草数据采集系统（实习 · 2021.09 – 2022.08）", bold=True)
    add_body(doc, "技术：Java、Spring Boot、MySQL、Vue")
    add_body(doc, "职责与成果", bold=True)
    for line in [
        "参与需求分析、原型设计、文档编写",
        "完成牌号管理、指标管理、报表导出等模块",
        "使用哨兵监控思维比对前端数组变化；解决 HashMap 并发问题",
        "项目成功上线",
    ]:
        add_bullet(doc, line)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"已生成: {OUTPUT}")


if __name__ == "__main__":
    build()
